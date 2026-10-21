"""Turning uploaded files and shared links into text. No network: DNS and HTTP are faked."""

from __future__ import annotations

import gzip
import io
import socket
import zipfile
import zlib
from collections.abc import Callable, Iterator

import docx
import httpcore
import httpx
import pypdf
import pytest
from docx.oxml import parse_xml

from agi_agent.knowledge import extract as ex
from agi_agent.knowledge.extract import (
    MAX_FILE_BYTES,
    MAX_TEXT_CHARS,
    USER_AGENT,
    ExtractError,
    classify_url,
    extract_file,
    extract_url,
)

PUBLIC_IP = "93.184.216.34"


# ---------------------------------------------------------------------------
# Helpers


def make_pdf(lines: list[str]) -> bytes:
    """A tiny valid one-page PDF with real text, built by hand."""
    content = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    stream = content.encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


RESUME_LINES = ["Shashank K - Applied AI engineer", "Builds document intelligence at PebbleRoad"]


def encrypt_pdf(data: bytes, user_password: str) -> bytes:
    writer = pypdf.PdfWriter(clone_from=pypdf.PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def save_docx(document: docx.document.Document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def resume_docx() -> bytes:
    """A Word file with the awkward bits real resume templates use."""
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "shashank@example.com | Singapore"
    document.add_heading("Shashank K", level=1)
    document.add_paragraph("Applied   AI engineer.")

    skills = document.add_table(rows=2, cols=2)
    for (row, col), value in {(0, 0): "Skill", (0, 1): "Level", (1, 0): "Python", (1, 1): "Expert"}.items():
        skills.cell(row, col).text = value

    merged = document.add_table(rows=2, cols=3)
    merged.cell(0, 0).text = "Merged once"
    merged.cell(0, 0).merge(merged.cell(0, 1))
    merged.cell(0, 2).text = "Tall cell"
    merged.cell(0, 2).merge(merged.cell(1, 2))
    merged.cell(1, 0).text = "Bottom left"

    layout = document.add_table(rows=1, cols=2)
    layout.cell(0, 0).text = "Experience"
    layout.cell(0, 0).add_paragraph("PebbleRoad, 2020-now")
    layout.cell(0, 1).text = "Education"

    anchor = document.add_paragraph("See box:")
    anchor._p.append(
        parse_xml(
            f"<w:r {NS}><mc:AlternateContent>"
            "<mc:Choice Requires='wps'><w:drawing><wps:txbx><w:txbxContent>"
            "<w:p><w:r><w:t>Boxed text</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:drawing></mc:Choice>"
            "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>"
            "<w:p><w:r><w:t>Boxed text</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
            "</mc:AlternateContent></w:r>"
        )
    )
    anchor._p.addnext(
        parse_xml(
            f"<w:sdt {NS}><w:sdtPr/><w:sdtContent>"
            "<w:p><w:r><w:t>Inside a content control</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        )
    )
    inline = document.add_paragraph("Phone: ")
    inline._p.append(parse_xml(f"<w:sdt {NS}><w:sdtContent><w:r><w:t>+65 1234 5678</w:t></w:r></w:sdtContent></w:sdt>"))
    return save_docx(document)


def html_response(
    body: str | bytes, status: int = 200, content_type: str = "text/html; charset=utf-8"
) -> httpx.Response:
    content = body.encode() if isinstance(body, str) else body
    return httpx.Response(status, headers={"Content-Type": content_type}, content=content)


Handler = Callable[[httpx.Request], httpx.Response]


def mock_client(handler: Handler, seen: list[httpx.Request] | None = None) -> httpx.Client:
    def record(request: httpx.Request) -> httpx.Response:
        if seen is not None:
            seen.append(request)
        return handler(request)

    return httpx.Client(transport=httpx.MockTransport(record))


def chunks(data: bytes, size: int = 16_384) -> Iterator[bytes]:
    for start in range(0, len(data), size):
        yield data[start : start + size]


@pytest.fixture(autouse=True)
def dns(monkeypatch: pytest.MonkeyPatch) -> dict[str, list[str]]:
    """Fake DNS so no test touches the network; tests may add names."""
    table = {
        "example.com": [PUBLIC_IP],
        "www.example.com": [PUBLIC_IP],
        "github.com": ["140.82.112.3"],
        "api.github.com": ["140.82.112.6"],
        "v6.example": ["2606:4700::1111"],
        "internal.example": ["10.1.2.3"],
        "mixed.example": [PUBLIC_IP, "192.168.1.10"],
        "mapped.example": ["::ffff:10.0.0.1"],
        "metadata.example": ["169.254.169.254"],
        "cgnat.example": ["100.64.0.1"],
        "loopback6.example": ["::1"],
    }

    def fake_getaddrinfo(host, port, family=0, type=0, proto=0, flags=0):  # noqa: A002 - mirrors socket's signature
        if host not in table:
            raise socket.gaierror(socket.EAI_NONAME, "Name or service not known")
        return [
            (socket.AF_INET6, socket.SOCK_STREAM, 6, "", (a, port, 0, 0))
            if ":" in a
            else (socket.AF_INET, socket.SOCK_STREAM, 6, "", (a, port))
            for a in table[host]
        ]

    monkeypatch.setattr(socket, "getaddrinfo", fake_getaddrinfo)
    return table


# ---------------------------------------------------------------------------
# Files


def test_pdf_text_is_extracted() -> None:
    result = extract_file("CV.pdf", make_pdf(RESUME_LINES))
    assert result.kind == "pdf"
    assert result.title == result.source == "CV.pdf"
    assert result.text == "\n".join(RESUME_LINES)
    assert not result.truncated and result.warnings == ()


def test_pdf_without_text_is_flagged_as_scanned() -> None:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    result = extract_file("scan.pdf", out.getvalue())
    assert result.text == ""
    assert any("looks scanned" in w for w in result.warnings)


def test_password_protected_pdf_is_refused() -> None:
    with pytest.raises(ExtractError, match="This PDF is password-protected"):
        extract_file("locked.pdf", encrypt_pdf(make_pdf(RESUME_LINES), user_password="hunter2"))


def test_pdf_encrypted_with_an_empty_password_still_opens() -> None:
    result = extract_file("restricted.pdf", encrypt_pdf(make_pdf(RESUME_LINES), user_password=""))
    assert "PebbleRoad" in result.text


@pytest.mark.parametrize("data", [b"%PDF-1.4\n" + b"garbage " * 20, b"%PDF-1.7\n"])
def test_damaged_pdf_is_refused(data: bytes) -> None:
    with pytest.raises(ExtractError, match="couldn't read this PDF"):
        extract_file("broken.pdf", data)


def test_docx_reads_paragraphs_tables_headers_text_boxes_and_content_controls() -> None:
    result = extract_file("resume.docx", resume_docx())
    text = result.text
    assert result.kind == "docx" and result.warnings == ()
    assert text.startswith("shashank@example.com | Singapore")
    assert "Shashank K\nApplied AI engineer." in text
    assert "Skill | Level\nPython | Expert" in text
    assert text.count("Merged once") == 1 and text.count("Tall cell") == 1
    assert "Merged once | Tall cell\nBottom left" in text
    assert "Experience\nPebbleRoad, 2020-now\n\nEducation" in text
    assert text.count("Boxed text") == 1
    assert "Inside a content control" in text
    assert "Phone: +65 1234 5678" in text


def test_docx_that_unpacks_too_large_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ex, "MAX_DOCX_UNPACKED_BYTES", 100)
    with pytest.raises(ExtractError, match="too large"):
        extract_file("bomb.docx", resume_docx())


def test_text_is_normalised() -> None:
    raw = "  Hello \t  world  \r\n\r\n\r\n\r\nSecond\x00 para\u00a0here  \n\x0cThird\ufeff \ufb01le\n\n"
    result = extract_file("notes.txt", raw.encode())
    assert result.kind == "text"
    assert result.text == "Hello world\n\nSecond para here\n\nThird file"


def test_text_decoding_handles_bad_bytes_and_utf16() -> None:
    assert extract_file("a.txt", b"caf\xe9 ok").text == "caf\ufffd ok"
    assert extract_file("b.md", "# Héllo".encode("utf-16")).text == "# Héllo"
    assert extract_file("NOTES.MD", b"\xef\xbb\xbfhi").text == "hi"


def test_long_text_is_truncated_with_a_note() -> None:
    result = extract_file("long.txt", ("word " * 20_000).encode())
    assert result.truncated
    assert len(result.text) <= MAX_TEXT_CHARS
    assert result.text.endswith("word")
    assert any("only the first part" in w for w in result.warnings)


def test_file_name_is_cleaned() -> None:
    assert extract_file("C:\\Users\\me\\My   Resume.txt", b"hi").title == "My Resume.txt"
    assert extract_file("../../etc/notes.md", b"hi").source == "notes.md"


def test_empty_text_file_gets_a_note() -> None:
    result = extract_file("blank.txt", b" \n\t\n")
    assert result.text == "" and result.warnings == ("No text was found in this file.",)


def _zip_without_word_document() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("hello.txt", "hi")
    return out.getvalue()


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("tool.exe", b"MZ\x90\x00", "isn't supported"),
        ("resume.doc", b"\xd0\xcf\x11\xe0", "save it as .docx"),
        ("noext", b"hello", "isn't supported"),
        ("resume.pdf", b"", "empty"),
        ("big.txt", b"a" * (MAX_FILE_BYTES + 1), "too big"),
        ("resume.pdf", b"just some text", "isn't a real PDF"),
        ("resume.pdf", b"PK\x03\x04 a zip", "isn't a real PDF"),
        ("resume.docx", b"just some text", "isn't a real Word document"),
        ("resume.docx", _zip_without_word_document(), "isn't a real Word document"),
        ("resume.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "older Word"),
        ("notes.txt", make_pdf(["hi"]), "doesn't look like plain text"),
        ("notes.md", b"PK\x03\x04" + b"\x00" * 64, "doesn't look like plain text"),
        ("notes.txt", b"\x00\x01\x02\x03" * 64, "doesn't look like plain text"),
    ],
    ids=lambda value: value if isinstance(value, str) else "data",
)
def test_bad_files_are_refused(filename: str, data: bytes, message: str) -> None:
    with pytest.raises(ExtractError, match=message):
        extract_file(filename, data)


# ---------------------------------------------------------------------------
# Links: classification and refusals


@pytest.mark.parametrize(
    ("url", "kind"),
    [
        ("https://www.linkedin.com/in/someone", "linkedin"),
        ("linkedin.com/in/someone", "linkedin"),
        ("https://in.linkedin.com/in/someone", "linkedin"),
        ("https://lnkd.in/abc", "linkedin"),
        ("https://github.com/octo", "github"),
        ("https://www.github.com/octo/ladder", "github"),
        ("https://x.com/someone", "x"),
        ("https://twitter.com/someone", "x"),
        ("https://mobile.twitter.com/someone", "x"),
        ("https://example.com/about", "web"),
        ("https://gist.github.com/octo", "web"),
        ("https://notlinkedin.com/in/x", "web"),
        ("https://linkedin.com.evil.example/", "web"),
        ("https://twitter.com.evil.example/", "web"),
        ("not a link at all", "web"),
    ],
)
def test_classify_url(url: str, kind: str) -> None:
    assert classify_url(url) == kind


@pytest.mark.parametrize(
    "url",
    ["https://www.linkedin.com/in/someone", "linkedin.com/in/someone", "https://lnkd.in/abc"],
)
def test_linkedin_is_never_fetched(url: str) -> None:
    seen: list[httpx.Request] = []
    with pytest.raises(ExtractError) as caught:
        extract_url(url, client=mock_client(lambda r: html_response("hi"), seen))
    assert str(caught.value) == (
        "LinkedIn doesn't let apps read profiles. Paste your About and Experience sections, "
        "or upload your LinkedIn profile PDF (More → Save to PDF on your profile)."
    )
    assert seen == []


@pytest.mark.parametrize("url", ["https://x.com/someone", "twitter.com/someone", "https://mobile.twitter.com/a"])
def test_x_is_never_fetched(url: str) -> None:
    seen: list[httpx.Request] = []
    with pytest.raises(ExtractError, match="Paste your bio"):
        extract_url(url, client=mock_client(lambda r: html_response("hi"), seen))
    assert seen == []


@pytest.mark.parametrize(
    ("url", "message"),
    [
        ("http://127.0.0.1/", "private or local"),
        ("http://10.0.0.5/admin", "private or local"),
        ("http://169.254.169.254/latest/meta-data/", "private or local"),
        ("http://localhost/", "private or local"),
        ("http://LOCALHOST./", "private or local"),
        ("http://admin.localhost/", "private or local"),
        ("http://[::1]/", "private or local"),
        ("http://[::ffff:127.0.0.1]/", "private or local"),
        ("http://[fe80::1]/", "private or local"),
        ("http://0.0.0.0/", "private or local"),
        ("http://224.0.0.1/", "private or local"),
        ("http://2130706433/", "private or local"),
        ("http://0x7f.1/", "private or local"),
        ("http://127.1/", "private or local"),
        ("http://internal.example/", "private or local"),
        ("http://mixed.example/", "private or local"),
        ("http://mapped.example/", "private or local"),
        ("http://metadata.example/", "private or local"),
        ("http://cgnat.example/", "private or local"),
        ("http://loopback6.example/", "private or local"),
        ("http://user:pass@example.com/", "username or password"),
        ("http://user@example.com/", "username or password"),
        ("http://example.com:22/", "standard web ports"),
        ("https://example.com:8443/", "standard web ports"),
        ("localhost:8080", "standard web ports"),
        ("file:///etc/passwd", "Only http"),
        ("ftp://example.com/file", "Only http"),
        ("gopher://example.com/", "Only http"),
        ("javascript:alert(1)", "Only http"),
        ("data:text/html,hi", "Only http"),
        ("http://nowhere.example/", "couldn't find that website"),
        ("https://example.com/" + "a" * 2100, "too long"),
        ("http:///no-host", "valid web address"),
        ("http://exa\nmple.com/", "valid web address"),
        ("", "enter a web address"),
    ],
)
def test_unsafe_links_are_refused_before_any_request(url: str, message: str) -> None:
    seen: list[httpx.Request] = []
    with pytest.raises(ExtractError, match=message):
        extract_url(url, client=mock_client(lambda r: html_response("hi"), seen))
    assert seen == []


@pytest.mark.parametrize(
    ("location", "message"),
    [
        ("http://169.254.169.254/latest/meta-data/", "private or local"),
        ("http://localhost:80/", "private or local"),
        ("http://internal.example/", "private or local"),
        ("http://[::1]/", "private or local"),
        ("file:///etc/passwd", "Only http"),
        ("http://example.com:22/", "standard web ports"),
        ("http://user:pw@example.com/", "username or password"),
        ("https://www.linkedin.com/in/someone", "LinkedIn doesn't let apps read profiles"),
    ],
)
def test_redirects_are_checked_at_every_hop(location: str, message: str) -> None:
    seen: list[httpx.Request] = []
    client = mock_client(lambda r: httpx.Response(302, headers={"Location": location}), seen)
    with pytest.raises(ExtractError, match=message):
        extract_url("https://example.com/me", client=client)
    assert [str(r.url) for r in seen] == ["https://example.com/me"]


def test_relative_redirects_are_followed() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/start":
            return httpx.Response(301, headers={"Location": "/about"})
        return html_response("<html><head><title>About</title></head><body><p>Hello there, I make things.</p></body>")

    result = extract_url("example.com/start", client=mock_client(handler))
    assert result.source == "https://example.com/about"
    assert "Hello there, I make things." in result.text


def test_too_many_redirects_stop() -> None:
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"Location": f"/hop{len(seen)}"})

    with pytest.raises(ExtractError, match="redirected too many times"):
        extract_url("https://example.com/", client=mock_client(handler, seen))
    assert len(seen) == ex.MAX_REDIRECTS + 1


def test_five_redirects_are_allowed() -> None:
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if len(seen) <= ex.MAX_REDIRECTS:
            return httpx.Response(307, headers={"Location": f"https://www.example.com/hop{len(seen)}"})
        return html_response("<p>Finally here.</p>", content_type="text/html")

    result = extract_url("https://example.com/", client=mock_client(handler, seen))
    assert result.text == "Finally here."
    assert result.source == f"https://www.example.com/hop{ex.MAX_REDIRECTS}"


def test_default_client_rechecks_the_address_at_connect_time(monkeypatch: pytest.MonkeyPatch) -> None:
    """DNS rebinding: public when checked, private a moment later. The connection must not happen."""
    answers = iter([PUBLIC_IP, "127.0.0.1"])

    def flipping(host, port, family=0, type=0, proto=0, flags=0):  # noqa: A002
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (next(answers), port))]

    connects: list[object] = []

    def no_connect(address, *args, **kwargs):
        connects.append(address)
        raise OSError("tests never open sockets")

    monkeypatch.setattr(socket, "getaddrinfo", flipping)
    monkeypatch.setattr(socket, "create_connection", no_connect)
    with pytest.raises(ExtractError, match="private or local"):
        extract_url("http://rebind.example/")
    assert connects == []


def test_guarded_backend_connects_to_the_checked_address(monkeypatch: pytest.MonkeyPatch) -> None:
    connects: list[object] = []

    def no_connect(address, *args, **kwargs):
        connects.append(address)
        raise OSError("tests never open sockets")

    monkeypatch.setattr(socket, "create_connection", no_connect)
    with pytest.raises(httpcore.ConnectError):
        ex._GuardedBackend().connect_tcp("example.com", 443)
    assert connects == [(PUBLIC_IP, 443)]
    with pytest.raises(ExtractError, match="private or local"):
        ex._GuardedBackend().connect_tcp("internal.example", 443)
    assert len(connects) == 1


# ---------------------------------------------------------------------------
# Links: GitHub

PROFILE = {
    "login": "octo",
    "name": "Octo Cat",
    "type": "User",
    "bio": "Builds\r\n  small   tools",
    "company": "@acme",
    "location": "Singapore",
    "blog": "https://octo.example",
    "twitter_username": "octocat",
    "public_repos": 12,
    "followers": 34,
    "created_at": "2011-01-25T18:44:36Z",
    "html_url": "https://github.com/octo",
}
REPOS = [
    {"name": "forked-thing", "fork": True, "description": "Someone else's"},
    {"name": "ladder", "fork": False, "description": "A tiny ladder", "language": "Python", "stargazers_count": 3},
    {"name": "quiet", "fork": False, "description": None, "language": None, "stargazers_count": 0},
]


def github_handler(
    profile: httpx.Response | None = None, repos: httpx.Response | None = None
) -> Callable[[httpx.Request], httpx.Response]:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.host == "api.github.com"
        if request.url.path == "/users/octo":
            return profile if profile is not None else httpx.Response(200, json=PROFILE)
        if request.url.path == "/users/octo/repos":
            assert request.url.params["sort"] == "pushed"
            return repos if repos is not None else httpx.Response(200, json=REPOS)
        return httpx.Response(404, json={"message": "Not Found"})

    return handler


def test_github_profile_is_read_through_the_api() -> None:
    seen: list[httpx.Request] = []
    result = extract_url("https://github.com/octo/", client=mock_client(github_handler(), seen))
    assert result.kind == "github"
    assert result.source == "https://github.com/octo"
    assert result.title == "Octo Cat (@octo) on GitHub"
    assert result.text.splitlines() == [
        "GitHub profile: Octo Cat (@octo)",
        "Bio: Builds small tools",
        "Company: @acme",
        "Location: Singapore",
        "Website: https://octo.example",
        "X (Twitter): @octocat",
        "12 public repositories; 34 followers",
        "On GitHub since 2011",
        "",
        "Recently active repositories:",
        "- ladder: A tiny ladder (Python, 3 stars)",
        "- quiet",
    ]
    assert all(r.headers["user-agent"] == USER_AGENT for r in seen)
    assert seen[0].headers["accept"] == "application/vnd.github+json"


def test_github_unknown_user() -> None:
    with pytest.raises(ExtractError, match="no GitHub account called"):
        extract_url("github.com/octo", client=mock_client(github_handler(profile=httpx.Response(404, json={}))))


def test_github_rate_limit() -> None:
    limited = httpx.Response(403, json={"message": "API rate limit exceeded"})
    with pytest.raises(ExtractError, match="limiting requests"):
        extract_url("github.com/octo", client=mock_client(github_handler(profile=limited)))


def test_github_profile_survives_a_repo_list_failure() -> None:
    result = extract_url("github.com/octo", client=mock_client(github_handler(repos=httpx.Response(500))))
    assert "Bio: Builds small tools" in result.text
    assert "Recently active" not in result.text
    assert any("repositories couldn't be loaded" in w for w in result.warnings)


@pytest.mark.parametrize("url", ["https://github.com/octo/ladder", "https://github.com/features"])
def test_other_github_pages_are_plain_web_pages(url: str) -> None:
    seen: list[httpx.Request] = []
    page = "<html><head><title>octo/ladder</title></head><body><p>A tiny ladder for climbing.</p></body></html>"
    result = extract_url(url, client=mock_client(lambda r: html_response(page), seen))
    assert result.kind == "web"
    assert [r.url.host for r in seen] == ["github.com"]
    assert "A tiny ladder for climbing." in result.text


# ---------------------------------------------------------------------------
# Links: web pages

ARTICLE = """<html><head><title>About me | Site</title><script>var x = "SECRET_SCRIPT";</script>
<style>.a{color:red}</style></head><body>
<header><nav><a href="/">Home</a><a href="/about">About</a><a href="/contact">Contact</a></nav></header>
<main><article><h1>About me</h1>
<p>I am a software engineer based in Singapore with ten years of experience building data products and search
systems for large organisations, and I care a lot about careful evaluation.</p>
<p>Before that I worked on document intelligence, retrieval and evaluation tooling, and I like writing about
applied machine learning in practice, especially where it meets messy real-world documents.</p>
<table><tr><th>Year</th><th>Role</th></tr><tr><td>2020</td><td>Lead engineer</td></tr></table>
</article></main><footer>Copyright 2026. All rights reserved.</footer></body></html>"""


def test_html_page_is_turned_into_text() -> None:
    seen: list[httpx.Request] = []
    result = extract_url("https://example.com/about", client=mock_client(lambda r: html_response(ARTICLE), seen))
    assert result.kind == "web"
    assert result.source == "https://example.com/about"
    assert "About me" in result.title
    assert "software engineer based in Singapore" in result.text
    assert "2020" in result.text and "Lead engineer" in result.text
    assert "SECRET_SCRIPT" not in result.text and "color:red" not in result.text
    assert "<" not in result.text
    assert seen[0].headers["user-agent"] == USER_AGENT
    assert "text/html" in seen[0].headers["accept"]


def test_short_homepage_keeps_the_owners_name() -> None:
    page = """<html><head><title>Shashank K</title></head><body><nav><a href="/">Home</a></nav>
    <h1>Hi, I'm Shashank.</h1><p>I build applied AI systems at PebbleRoad.</p></body></html>"""
    result = extract_url("https://example.com/", client=mock_client(lambda r: html_response(page)))
    assert "Hi, I'm Shashank." in result.text
    assert "I build applied AI systems at PebbleRoad." in result.text


def test_tag_stripping_fallback(monkeypatch: pytest.MonkeyPatch) -> None:
    def broken(*args, **kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(ex.trafilatura, "extract", lambda *a, **k: None)
    monkeypatch.setattr(ex.trafilatura, "extract_metadata", broken)
    page = (
        "<html><head><title>Fallback &amp; friends</title><script>SECRET()</script></head>"
        "<body><div>Hello from the fallback</div><p>Second&nbsp;line<br>third</p><svg><text>icon</text></svg></body>"
    )
    result = extract_url("https://example.com/", client=mock_client(lambda r: html_response(page)))
    assert result.title == "Fallback & friends"
    assert result.text == "Hello from the fallback\n\nSecond line\nthird"


def test_javascript_only_page_gets_a_note() -> None:
    page = '<html><head><title>App</title></head><body><div id="root"></div><script src="/app.js"></script></body>'
    result = extract_url("https://example.com/", client=mock_client(lambda r: html_response(page)))
    assert result.text == ""
    assert any("JavaScript" in w for w in result.warnings)


def test_plain_text_and_xhtml_are_accepted() -> None:
    text_page = html_response("caf\xe9 notes".encode("latin-1"), content_type="text/plain; charset=ISO-8859-1")
    result = extract_url("https://example.com/files/notes.txt", client=mock_client(lambda r: text_page))
    assert (result.text, result.title, result.kind) == ("café notes", "notes.txt", "web")

    xhtml = html_response("<html><body><p>Hello XHTML world.</p></body></html>", content_type="application/xhtml+xml")
    assert "Hello XHTML world." in extract_url("https://example.com/", client=mock_client(lambda r: xhtml)).text


@pytest.mark.parametrize(
    ("content_type", "message"),
    [
        ("application/pdf", "points to a PDF"),
        ("image/png", "image"),
        ("application/json", "isn't a web page we can read"),
        ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "document file"),
        (None, "didn't say what kind"),
    ],
)
def test_other_content_types_are_refused(content_type: str | None, message: str) -> None:
    headers = {"Content-Type": content_type} if content_type else {}
    client = mock_client(lambda r: httpx.Response(200, headers=headers, content=b"%PDF-1.4 binary"))
    with pytest.raises(ExtractError, match=message):
        extract_url("https://example.com/file", client=client)


def test_body_is_streamed_and_cut_at_the_size_limit() -> None:
    pulled: list[int] = []

    def endless() -> Iterator[bytes]:
        for index in range(200):  # 12.5 MB on offer
            pulled.append(index)
            yield b"lorem ipsum " * 5461  # ~64 KB

    client = mock_client(lambda r: httpx.Response(200, headers={"Content-Type": "text/plain"}, content=endless()))
    result = extract_url("https://example.com/huge.txt", client=client)
    assert len(pulled) <= ex.MAX_RESPONSE_BYTES // 65_532 + 2
    assert result.truncated and len(result.text) <= MAX_TEXT_CHARS
    assert any("very large" in w for w in result.warnings)


def test_compressed_body_is_capped_after_decompression() -> None:
    bomb = gzip.compress(b"a" * (20 * 1024 * 1024))
    response = httpx.Response(
        200, headers={"Content-Type": "text/plain", "Content-Encoding": "gzip"}, content=chunks(bomb)
    )
    result = extract_url("https://example.com/bomb", client=mock_client(lambda r: response))
    assert result.truncated
    assert set(result.text) == {"a"} and len(result.text) <= MAX_TEXT_CHARS


@pytest.mark.parametrize(
    ("encoding", "compress"),
    [
        ("gzip", gzip.compress),
        ("deflate", zlib.compress),
        ("deflate", lambda data: (lambda c: c.compress(data) + c.flush())(zlib.compressobj(wbits=-15))),
    ],
)
def test_compressed_pages_are_decoded(encoding: str, compress: Callable[[bytes], bytes]) -> None:
    body = compress(b"<html><body><p>Squeezed but readable.</p></body></html>")

    def handler(request: httpx.Request) -> httpx.Response:
        assert "gzip" in request.headers["accept-encoding"]
        headers = {"Content-Type": "text/html", "Content-Encoding": encoding}
        return httpx.Response(200, headers=headers, content=chunks(body, 7))

    assert extract_url("https://example.com/", client=mock_client(handler)).text == "Squeezed but readable."


def test_unknown_compression_is_refused() -> None:
    response = httpx.Response(
        200, headers={"Content-Type": "text/html", "Content-Encoding": "br"}, content=iter([b"x"])
    )
    with pytest.raises(ExtractError, match="compressed format"):
        extract_url("https://example.com/", client=mock_client(lambda r: response))


@pytest.mark.parametrize(
    ("status", "message"),
    [
        (404, "wasn't found"),
        (410, "wasn't found"),
        (403, "private or blocks apps"),
        (429, "limiting"),
        (503, "problems"),
    ],
)
def test_error_statuses_have_plain_messages(status: int, message: str) -> None:
    client = mock_client(lambda r: html_response("<p>SECRET page body</p>", status=status))
    with pytest.raises(ExtractError, match=message) as caught:
        extract_url("https://example.com/", client=client)
    assert "SECRET" not in str(caught.value)


@pytest.mark.parametrize(
    ("error", "message"),
    [(httpx.ReadTimeout, "took too long"), (httpx.ConnectError, "couldn't connect")],
)
def test_network_failures_have_plain_messages(error: type[httpx.TransportError], message: str) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise error("simulated", request=request)

    with pytest.raises(ExtractError, match=message):
        extract_url("https://example.com/", client=mock_client(handler))


def test_slow_drip_body_hits_the_overall_deadline(monkeypatch: pytest.MonkeyPatch) -> None:
    clock = [0.0]

    class FakeTime:
        @staticmethod
        def monotonic() -> float:
            return clock[0]

    def drip() -> Iterator[bytes]:
        while True:
            clock[0] += 5
            yield b"a"

    monkeypatch.setattr(ex, "time", FakeTime)
    client = mock_client(lambda r: httpx.Response(200, headers={"Content-Type": "text/plain"}, content=drip()))
    with pytest.raises(ExtractError, match="took too long"):
        extract_url("https://example.com/", client=client)


def test_ipv6_public_hosts_are_allowed() -> None:
    result = extract_url("https://v6.example/", client=mock_client(lambda r: html_response("<p>Over IPv6.</p>")))
    assert result.text == "Over IPv6."
